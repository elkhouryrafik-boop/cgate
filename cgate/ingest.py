"""
cgate.ingest — source files to a canonical, reading-ordered Document.

Why this module is the most important one in the tool
-----------------------------------------------------
A tender can carry its General Terms and Conditions on
several pages in a two-column layout. A naive text extraction interleaves the two
columns line by line and produces text in which no sentence survives intact. A
compliance reviewer reading that output would conclude the General Conditions
were not attached at all — which is precisely the wrong finding, and it is a
finding about the *presence of a document*, not about any subtlety of law.

Everything downstream inherits ingest's mistakes and cannot detect them. So the
column reconstruction here is not a nicety; it is the difference between a tool
that reads the contract and a tool that hallucinates about one.

Canonical text
--------------
`Block.text` is normalised (NFKC, ligatures folded, quotes straightened,
de-hyphenated across line breaks, whitespace collapsed). All span offsets index
into that normalised text, and the grounding gate verifies against it. The raw
bytes remain addressable through `doc_id` + page, which is what an auditor needs.

Failure contract
----------------
An unreadable source produces a structured error with `retryable`, never an
empty Document that looks like a document containing nothing. (CCA practice P6.)
"""

from __future__ import annotations

import os
import re
import unicodedata

from .model import Block, Document, sha256_file

# ---------------------------------------------------------------------------
# Layout constants
# ---------------------------------------------------------------------------

MIN_GUTTER_PT = 6.0          # narrower than this is inter-word space, not a column gap
FULL_WIDTH_RATIO = 0.72      # a block this wide spans columns and acts as a separator
HEADER_BAND = 0.06           # top 6% of the page
FOOTER_BAND = 0.94           # bottom 6% of the page
MIN_BLOCK_CHARS = 2
LOW_SPACE_RATIO = 0.05       # below this, word boundaries were probably lost

_LIGATURES = {
    "ﬀ": "ff", "ﬁ": "fi", "ﬂ": "fl",
    "ﬃ": "ffi", "ﬄ": "ffl", "ﬅ": "st", "ﬆ": "st",
}
_QUOTES = {
    "‘": "'", "’": "'", "‚": "'", "‛": "'",
    "“": '"', "”": '"', "„": '"',
    "–": "-", "—": "-", "−": "-", "­": "",
    " ": " ", " ": " ", " ": " ", "\ufeff": "",
}


def normalise(text: str) -> str:
    """Fold the source's typography without destroying its words."""
    text = unicodedata.normalize("NFKC", text)
    for src, dst in {**_LIGATURES, **_QUOTES}.items():
        text = text.replace(src, dst)
    # De-hyphenate across a line break: "sub-\ncontractor" -> "subcontractor".
    # Only when the continuation is lower-case, so "Anti-\nBribery" survives.
    text = re.sub(r"(\w)-\s*\n\s*([a-z])", r"\1\2", text)
    text = re.sub(r"[ \t]*\n[ \t]*", " ", text)
    text = re.sub(r"[ \t]{2,}", " ", text)
    return text.strip()


# ---------------------------------------------------------------------------
# Column geometry
# ---------------------------------------------------------------------------


def _find_gutters(boxes: list[tuple[float, float]], x0: float, x1: float) -> list[float]:
    """Vertical bands of the page that body text essentially does not occupy.

    Implemented as a tolerant horizontal projection profile rather than a
    strict gap sweep. The strict version is what most extractors do and it is
    fragile for exactly one reason: a single merged or mis-boxed block that
    straddles the gutter collapses the two columns into one, and the whole
    page then reads as interleaved nonsense. Measured on two-column general-conditions
    pages, one straddling block in 106 was enough to do it.

    So a gutter is a run of x where coverage falls to a small tolerance, not to
    zero, and it must additionally have real text mass on both sides.
    """
    if len(boxes) < 8:
        return []
    lo, hi = int(x0), int(x1) + 1
    if hi - lo < 40:
        return []

    width = hi - lo
    cover = [0] * width
    for left, right in boxes:
        a = max(0, int(left) - lo)
        b = min(width, int(right) - lo + 1)
        for i in range(a, b):
            cover[i] += 1

    # Threshold relative to how densely this page is set, not to the block
    # count. A dense two-column page still shows ~2-3 residual coverage in the
    # gutter from stragglers; a fixed tolerance either misses it on dense pages
    # or invents gutters on sparse ones.
    busy = sorted(c for c in cover if c > 0)
    typical = busy[len(busy) // 2] if busy else 0
    tol = max(1, int(0.10 * typical))
    margin = int(0.12 * width)

    gutters: list[float] = []
    run_start: int | None = None
    for i in range(width + 1):
        empty = i < width and cover[i] <= tol
        if empty and run_start is None:
            run_start = i
        elif not empty and run_start is not None:
            run_len = i - run_start
            mid = run_start + run_len / 2.0
            if run_len >= MIN_GUTTER_PT and margin < mid < width - margin:
                left_mass = sum(cover[:run_start])
                right_mass = sum(cover[i:])
                if left_mass > 0 and right_mass > 0 and \
                        min(left_mass, right_mass) >= 0.15 * (left_mass + right_mass):
                    gutters.append(lo + mid)
            run_start = None
    return gutters


def _column_of(left: float, right: float, gutters: list[float]) -> int:
    mid = (left + right) / 2.0
    col = 0
    for g in gutters:
        if mid > g:
            col += 1
    return col


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------


def _ingest_pdf(path: str, doc_id: str) -> Document:
    import pymupdf

    src = pymupdf.open(path)
    doc = Document(doc_id=doc_id, path=path, filetype="pdf", n_pages=src.page_count)
    order = 0
    body_sizes: list[float] = []

    try:
        # First pass: establish the body font size, used for heading detection.
        for pno in range(src.page_count):
            for blk in src[pno].get_text("dict")["blocks"]:
                for line in blk.get("lines", []):
                    for sp in line.get("spans", []):
                        if len(sp.get("text", "").strip()) > 3:
                            body_sizes.append(round(sp["size"], 1))
        body_size = _mode(body_sizes) if body_sizes else 10.0

        for pno in range(src.page_count):
            page = src[pno]
            pw, ph = page.rect.width, page.rect.height
            raw = [b for b in page.get_text("dict")["blocks"] if b.get("lines")]

            entries = []
            for blk in raw:
                text_parts, sizes, bold = [], [], 0
                nspans = 0
                for line in blk["lines"]:
                    line_txt = "".join(sp["text"] for sp in line["spans"])
                    text_parts.append(line_txt)
                    for sp in line["spans"]:
                        sizes.append(sp["size"])
                        nspans += 1
                        if sp.get("flags", 0) & 16:      # bold bit
                            bold += 1
                text = normalise("\n".join(text_parts))
                if len(text) < MIN_BLOCK_CHARS:
                    continue
                x0b, y0b, x1b, y1b = blk["bbox"]
                entries.append(
                    {
                        "text": text,
                        "bbox": (x0b, y0b, x1b, y1b),
                        "size": max(sizes) if sizes else body_size,
                        "bold_ratio": (bold / nspans) if nspans else 0.0,
                        "full_width": (x1b - x0b) >= FULL_WIDTH_RATIO * pw,
                    }
                )

            if not entries:
                continue

            text_x0 = min(e["bbox"][0] for e in entries)
            text_x1 = max(e["bbox"][2] for e in entries)
            narrow = [(e["bbox"][0], e["bbox"][2]) for e in entries if not e["full_width"]]
            gutters = _find_gutters(narrow, text_x0, text_x1)

            for e in entries:
                e["col"] = -1 if e["full_width"] else _column_of(
                    e["bbox"][0], e["bbox"][2], gutters
                )

            # Full-width blocks are horizontal separators. Read each band's
            # columns in full before moving to the next band.
            seps = sorted(e["bbox"][1] for e in entries if e["col"] == -1)
            def band_of(y: float) -> int:
                return sum(1 for s in seps if y >= s - 0.5)

            for e in entries:
                e["band"] = band_of(e["bbox"][1])

            entries.sort(
                key=lambda e: (
                    e["band"],
                    99 if e["col"] == -1 else e["col"],
                    round(e["bbox"][1], 1),
                    round(e["bbox"][0], 1),
                )
            )

            for i, e in enumerate(entries):
                y0b = e["bbox"][1]
                kind = "body"
                if y0b < HEADER_BAND * ph:
                    kind = "header"
                elif y0b > FOOTER_BAND * ph:
                    kind = "footer"
                elif _looks_like_heading(e, body_size):
                    kind = "heading"
                col = 0 if e["col"] == -1 else e["col"]
                doc.blocks.append(
                    Block(
                        block_id=f"p{pno + 1}.c{col}.b{i}",
                        page=pno + 1,
                        column=col,
                        order=order,
                        text=e["text"],
                        bbox=e["bbox"],
                        kind=kind,
                    )
                )
                order += 1

            if len(gutters) >= 1:
                doc.notes.append(
                    f"page {pno + 1}: {len(gutters) + 1} columns detected "
                    f"(gutters at x={[round(g) for g in gutters]})"
                )
    finally:
        src.close()

    doc.build_index()
    _post_checks(doc)
    return doc


def _mode(values: list[float]) -> float:
    counts: dict[float, int] = {}
    for v in values:
        counts[v] = counts.get(v, 0) + 1
    return max(counts, key=counts.get) if counts else 10.0


_NUMBER_HEAD = re.compile(
    r"^(?:\d+(?:\.\d+)*\.?|[A-Z]\.|\([a-z]\)|[ivxlc]+\.)\s+\S", re.IGNORECASE
)


def _looks_like_heading(e: dict, body_size: float) -> bool:
    text = e["text"]
    if len(text) > 140:
        return False
    larger = e["size"] >= body_size + 0.6
    boldish = e["bold_ratio"] >= 0.6
    titled = text.isupper() and len(text) > 3
    if larger or titled:
        return True
    if boldish and (len(text) < 90 or _NUMBER_HEAD.match(text)):
        return True
    return False


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------


def _ingest_docx(path: str, doc_id: str) -> Document:
    from docx import Document as Docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    src = Docx(path)
    doc = Document(doc_id=doc_id, path=path, filetype="docx", n_pages=0)
    order = 0

    body = src.element.body
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            para = Paragraph(child, src)
            text = normalise(para.text)
            if len(text) < MIN_BLOCK_CHARS:
                continue
            style = getattr(para.style, "name", "") or ""
            kind = "heading" if style.lower().startswith("heading") or (
                text.isupper() and len(text) < 120
            ) else "body"
            doc.blocks.append(
                Block(f"d.b{order}", 1, 0, order, text, (0, 0, 0, 0), kind)
            )
            order += 1
        elif tag == "tbl":
            table = Table(child, src)
            for r, row in enumerate(table.rows):
                cells = [normalise(c.text) for c in row.cells]
                text = " | ".join(c for c in cells if c)
                if len(text) < MIN_BLOCK_CHARS:
                    continue
                doc.blocks.append(
                    Block(f"d.t{order}.r{r}", 1, 0, order, text, (0, 0, 0, 0), "table")
                )
                order += 1

    doc.build_index()
    _post_checks(doc)
    return doc


# ---------------------------------------------------------------------------


def _post_checks(doc: Document) -> None:
    """Observations ingest is entitled to make about its own output."""
    bad_spacing = 0
    for b in doc.blocks:
        if len(b.text) > 40:
            ratio = b.text.count(" ") / len(b.text)
            if ratio < LOW_SPACE_RATIO:
                bad_spacing += 1
    if bad_spacing:
        doc.notes.append(
            f"{bad_spacing} block(s) show lost word spacing - the source font "
            f"encoding may not carry space glyphs; extraction is unreliable"
        )
        doc.errors.append(
            {
                "error_type": "word_spacing_lost",
                "message": f"{bad_spacing} blocks",
                "retryable": False,
                "blocking": bad_spacing > len(doc.blocks) * 0.1,
            }
        )
    if not doc.blocks:
        doc.errors.append(
            {
                "error_type": "no_text_extracted",
                "message": "document produced zero text blocks; it may be a scan",
                "retryable": False,
                "blocking": True,
            }
        )


def ingest(path: str) -> Document:
    """Read a PDF or DOCX into a reading-ordered Document."""
    if not os.path.exists(path):
        d = Document(doc_id="", path=path, filetype="", n_pages=0)
        d.errors.append(
            {"error_type": "not_found", "message": path,
             "retryable": False, "blocking": True}
        )
        return d

    doc_id = sha256_file(path)
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".pdf":
            return _ingest_pdf(path, doc_id)
        if ext == ".docx":
            return _ingest_docx(path, doc_id)
        if ext in (".txt", ".md"):
            with open(path, encoding="utf-8", errors="replace") as fh:
                raw = fh.read()
            d = Document(doc_id=doc_id, path=path, filetype="text", n_pages=1)
            for i, para in enumerate(re.split(r"\n\s*\n", raw)):
                t = normalise(para)
                if len(t) >= MIN_BLOCK_CHARS:
                    d.blocks.append(
                        Block(f"t.b{i}", 1, 0, i, t, (0, 0, 0, 0),
                              "heading" if t.isupper() and len(t) < 120 else "body")
                    )
            d.build_index()
            _post_checks(d)
            return d
    except Exception as exc:                                    # noqa: BLE001
        d = Document(doc_id=doc_id, path=path, filetype=ext, n_pages=0)
        d.errors.append(
            {"error_type": "extraction_failed", "message": f"{type(exc).__name__}: {exc}",
             "retryable": True, "blocking": True}
        )
        return d

    d = Document(doc_id=doc_id, path=path, filetype=ext, n_pages=0)
    d.errors.append(
        {"error_type": "unsupported_type", "message": ext,
         "retryable": False, "blocking": True}
    )
    return d
